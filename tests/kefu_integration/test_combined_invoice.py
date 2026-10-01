"""
Combined multi-warehouse 费用报告 (invoice-inventory plan, Phase 3).

Fees here come from storage-fee ledger rows seeded for two test-only
warehouse codes in a 2031 test month (nothing real can collide); every row
is removed by exact id. The real view_invoice schema (V38) is read from the
database. Real Postgres DB.
"""
import datetime
import io
from datetime import timedelta, timezone
from decimal import Decimal

import pytest
from openpyxl import load_workbook
from sqlalchemy import text

from database import SessionLocal

WH_A, WH_B = "CIA", "CIB"
MONTH = "2031-05"
FIXED = datetime.datetime(2031, 6, 1, tzinfo=timezone.utc)


@pytest.fixture
def db():
    session = SessionLocal()
    yield session
    session.rollback()
    session.close()


@pytest.fixture
def ledger(db):
    """CIA: 2 days at 4 pallets ($2.00/day); CIB: 1 day at 10 pallets ($5.00)."""
    from models.uchoice import UchoiceStorageFeeLedger

    rows = [
        UchoiceStorageFeeLedger(warehouse_code=WH_A, fee_date=datetime.date(2031, 5, 1), pallet_count=4, storage_fee=Decimal("2.00")),
        UchoiceStorageFeeLedger(warehouse_code=WH_A, fee_date=datetime.date(2031, 5, 2), pallet_count=4, storage_fee=Decimal("2.00")),
        UchoiceStorageFeeLedger(warehouse_code=WH_B, fee_date=datetime.date(2031, 5, 1), pallet_count=10, storage_fee=Decimal("5.00")),
    ]
    db.add_all(rows)
    db.commit()
    yield
    db.rollback()
    for row in rows:
        db.execute(text("delete from uchoice_storage_fee_ledger where ledger_id = :id"), {"id": row.ledger_id})
    db.commit()


def test_combined_invoice_adds_up_per_warehouse_totals(db, ledger):
    from core.uchoice_invoice import compute_combined_invoice, compute_invoice

    combined = compute_combined_invoice(db, [WH_B, WH_A, WH_B], MONTH)
    assert combined["warehouse_codes"] == [WH_A, WH_B]
    assert combined["storage_fee"] == Decimal("9.00")
    assert combined["total"] == sum((compute_invoice(db, c, MONTH)["total"] for c in (WH_A, WH_B)), Decimal("0"))
    assert combined["per_warehouse"][WH_A]["storage_fee"] == Decimal("4.00")


def test_combined_workbook_summary_by_warehouse(db, ledger):
    from core.uchoice_invoice_export import build_invoice_workbook

    wb = load_workbook(io.BytesIO(build_invoice_workbook(db, [WH_B, WH_A], MONTH, generated_at=FIXED)))
    ws = wb["Summary"]
    rows = [[c.value for c in row] for row in ws.iter_rows()]
    assert rows[0] == ["Warehouses", "CIA, CIB", None, None]
    header = rows.index(["Charge (USD)", WH_A, WH_B, "Total"])
    storage = rows[header + 4]
    assert storage[:3] == ["Storage fee", 4.0, 5.0]
    assert storage[3] == f"=SUM(B{header + 5}:C{header + 5})"
    total = rows[header + 5]
    assert total[0] == "Total"
    assert total[1:] == [f"=SUM({col}{header + 2}:{col}{header + 5})" for col in "BCD"]
    assert ["Inventory (plt)", WH_A, WH_B, "Total"] in rows

    storage_rows = list(wb["Storage"].iter_rows(min_row=2, values_only=True))
    assert [r[0] for r in storage_rows] == [WH_A, WH_A, WH_B]   # grouped by warehouse


def test_single_warehouse_uses_the_same_layout(db, ledger):
    from core.uchoice_invoice_export import build_invoice_workbook

    wb = load_workbook(io.BytesIO(build_invoice_workbook(db, WH_A, MONTH, generated_at=FIXED)))
    rows = [[c.value for c in row] for row in wb["Summary"].iter_rows()]
    assert ["Charge (USD)", WH_A, "Total"] in rows


def test_multi_warehouse_workbook_is_deterministic(db, ledger):
    from core.uchoice_invoice_export import build_invoice_workbook

    first = build_invoice_workbook(db, [WH_A, WH_B], MONTH, generated_at=FIXED)
    second = build_invoice_workbook(db, [WH_B, WH_A], MONTH, generated_at=FIXED)
    assert first == second


def test_filename_lists_warehouses_sorted():
    from core.uchoice_invoice_export import invoice_filename
    assert invoice_filename(["NJ", "DE", "JFK"], "2026-09", None) == "invoice_DE-JFK-NJ_2026-09_2026-09.xlsx"
    assert invoice_filename("JFK", "2026-08", "2026-09") == "invoice_JFK_2026-08_2026-09.xlsx"


def test_v38_schema_no_longer_asks_for_a_warehouse(db):
    from core import kefu_turn_apply

    schema = db.execute(text("select input_schema from service_type where name = 'view_invoice'")).scalar()
    assert schema["required"] == ["start_month", "end_month"]
    assert schema["optional"] == ["warehouse_codes"]
    service = {"input_schema": schema}
    assert kefu_turn_apply._all_required_fields_present(service, {"start_month": MONTH, "end_month": MONTH})
    prompt = kefu_turn_apply._render_missing_fields(service, {})
    assert "仓库" not in prompt


def test_smart_robot_handler_defaults_to_all_accessible(db, ledger):
    from handlers.uchoice.queries import ComputeInvoiceHandler

    result = ComputeInvoiceHandler().handle({
        "role": "warehouseman", "warehouse_codes": [WH_B, WH_A], "group_id": None,
        "collected_fields": {"start_month": MONTH, "end_month": MONTH},
    }, {}, db)
    assert result["warehouse_codes"] == [WH_A, WH_B]
    assert result["storage_fee"] == "9.00"
    assert set(result["per_warehouse"]) == {WH_A, WH_B}
    assert result.get("download_url")


def test_kefu_turn_builds_one_combined_file(db, ledger, monkeypatch):
    from core import kefu_turn_apply
    from models.group import GroupConfig
    from models.request_log import RequestLog
    from models.service import ServiceType
    from models.session import ConversationSession
    from models.workflow import Workflow

    group = db.query(GroupConfig).order_by(GroupConfig.created_at).first()
    service_type = db.query(ServiceType).filter_by(name="view_invoice").one()
    workflow = db.query(Workflow).filter_by(name="view_invoice").one()
    service = {"name": "view_invoice", "service_type_id": str(service_type.service_type_id),
               "workflow_id": str(workflow.workflow_id), "requires_confirmation": False,
               "targets_existing_request": False}
    now = datetime.datetime.now(timezone.utc)
    session = ConversationSession(group_id=group.group_id, service_type_id=service_type.service_type_id,
                                  status="active", conversation_history=[], source_channel="kefu",
                                  collected_fields={"warehouse_codes": [WH_B, WH_A], "start_month": MONTH, "end_month": MONTH},
                                  expires_at=now + timedelta(minutes=30))
    log = RequestLog(group_id=group.group_id, service_type_id=service_type.service_type_id,
                     status="pending", raw_message="combined invoice test", source_channel="kefu")
    db.add_all([session, log])
    db.flush()
    session.request_log_id = log.log_id
    db.commit()
    _forbid_second_calculation(monkeypatch)
    try:
        context = {"result": {}, "request_log_id": str(log.log_id), "role": "admin", "warehouse_codes": None}
        kefu_turn_apply._workflow_steps(db, context, service, session)
        assert context["result"]["warehouse_codes"] == [WH_A, WH_B]
        [artifact] = context["_kefu_artifacts"]
        assert artifact["artifact"]["filename"] == f"invoice_CIA-CIB_{MONTH}_{MONTH}.xlsx"
    finally:
        db.rollback()
        db.execute(text("delete from conversation_session where session_id = :s"), {"s": session.session_id})
        db.execute(text("delete from request_log where log_id = :l"), {"l": log.log_id})
        db.commit()


def test_chat_reply_lists_warehouses_and_per_warehouse_totals():
    from core.message_sections import render_sections
    from core.result_message import _invoice_sections_builder

    lines = render_sections(_invoice_sections_builder({"result": {
        "warehouse_codes": ["DE", "JFK"], "start_month": "2026-09", "end_month": "2026-09",
        "total": "30.00", "opening_pallets": 93, "closing_pallets": 88,
        "per_warehouse": {"DE": {"total": "10.00"}, "JFK": {"total": "20.00"}},
    }}, None))
    assert lines[0] == "仓库：DE、JFK　范围：2026-09"
    assert "**各仓合计**" in lines
    assert "- DE：$10.00" in lines and "- JFK：$20.00" in lines


def _forbid_second_calculation(monkeypatch):
    """Codex code audit #4: the reply must reuse the workbook's own numbers.
    A separate compute_combined_invoice call is what let them diverge."""
    import core.uchoice_invoice as uchoice_invoice

    def _second_calculation(*args, **kwargs):
        raise AssertionError("invoice computed a second time for the reply")
    monkeypatch.setattr(uchoice_invoice, "compute_combined_invoice", _second_calculation)


def test_report_totals_match_its_own_sheets(db, ledger):
    from core.uchoice_invoice_export import build_invoice_report

    data, invoice = build_invoice_report(db, [WH_A, WH_B], MONTH, generated_at=FIXED)
    wb = load_workbook(io.BytesIO(data))
    summary = {r[0]: r[1:] for r in wb["Summary"].iter_rows(values_only=True) if r and r[0]}
    assert summary["Storage fee"][:2] == (4.0, 5.0)
    assert invoice["storage_fee"] == Decimal("9.00")
    inventory = list(wb["Inventory"].iter_rows(min_row=2, values_only=True))
    opening = sum(r[3] for r in inventory if r[0] in (WH_A, WH_B))
    assert invoice["opening_pallets"] == opening
    assert summary["Opening pallets"][:2] == (invoice["per_warehouse"][WH_A]["opening_pallets"],
                                              invoice["per_warehouse"][WH_B]["opening_pallets"])


def test_smart_robot_reply_comes_from_the_workbook_calculation(db, ledger, monkeypatch):
    from handlers.uchoice.queries import ComputeInvoiceHandler

    _forbid_second_calculation(monkeypatch)
    result = ComputeInvoiceHandler().handle({
        "role": "admin", "warehouse_codes": None, "group_id": None,
        "collected_fields": {"warehouse_codes": [WH_A, WH_B], "start_month": MONTH, "end_month": MONTH},
    }, {}, db)
    assert result["storage_fee"] == "9.00" and result.get("download_url")


def test_smart_robot_reply_still_works_if_the_workbook_fails(db, ledger, monkeypatch):
    import core.uchoice_invoice_export as export
    from handlers.uchoice.queries import ComputeInvoiceHandler

    def _broken(*args, **kwargs):
        raise RuntimeError("workbook failed")
    monkeypatch.setattr(export, "build_invoice_report", _broken)
    result = ComputeInvoiceHandler().handle({
        "role": "admin", "warehouse_codes": None, "group_id": None,
        "collected_fields": {"warehouse_codes": [WH_A], "start_month": MONTH, "end_month": MONTH},
    }, {}, db)
    assert result["storage_fee"] == "4.00" and "download_url" not in result


def test_totals_equal_detail_rows_when_a_row_lands_mid_build(db, ledger, monkeypatch):
    """Codex code audit round 2 #1: the fee rows are selected once, so a ledger
    row committed between the totals and the detail sheets is in neither."""
    import core.uchoice_invoice_export as export
    from models.uchoice import UchoiceStorageFeeLedger

    late = []
    real_balances = export.inventory_balances

    def _balances_then_concurrent_write(*args, **kwargs):
        # Runs after the fee rows are selected, before the sheets are written.
        other = SessionLocal()
        row = UchoiceStorageFeeLedger(warehouse_code=WH_A, fee_date=datetime.date(2031, 5, 3),
                                      pallet_count=4, storage_fee=Decimal("2.00"))
        other.add(row)
        other.commit()
        late.append(row.ledger_id)
        other.close()
        return real_balances(*args, **kwargs)

    monkeypatch.setattr(export, "inventory_balances", _balances_then_concurrent_write)
    try:
        data, invoice = export.build_invoice_report(db, [WH_A], MONTH, generated_at=FIXED)
        storage_rows = list(load_workbook(io.BytesIO(data))["Storage"].iter_rows(min_row=2, values_only=True))
        assert invoice["storage_fee"] == Decimal("4.00")
        assert sum(Decimal(str(r[3])) for r in storage_rows) == invoice["storage_fee"]
        assert len(storage_rows) == 2
    finally:
        db.rollback()
        for ledger_id in late:
            db.execute(text("delete from uchoice_storage_fee_ledger where ledger_id = :id"), {"id": ledger_id})
        db.commit()
