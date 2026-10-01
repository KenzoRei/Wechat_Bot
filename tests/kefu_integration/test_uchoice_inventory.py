"""
Invoice Inventory sheet (invoice-inventory plan, Phase 2 + 2b):
core/uchoice_inventory.inventory_balances and the workbook sheet built on it.

Stock changes are made through the real code paths
(core.uchoice_storage.apply_storage_delta / apply_loose_pick and
ApplyInbound/ApplyOutboundStorageHandler), then backdated so each lands in
a known spot relative to a 2031-03 test period. Warehouse codes are
test-only (not in VALID_WAREHOUSE_CODES), so nothing here can collide with
real data; cleanup deletes exactly those warehouses' rows.

Real Postgres DB.
"""
import datetime
import io

import pytest
from openpyxl import load_workbook
from sqlalchemy import text

from database import SessionLocal

WH = "INVTEST"
DEST = "INVDEST"
SKU = "s2"
SKU_B = "t1"
START = datetime.date(2031, 3, 1)
END_EXCLUSIVE = datetime.date(2031, 4, 1)
UTC = datetime.timezone.utc


def _cleanup(db):
    db.rollback()
    for table in ("uchoice_storage_txn", "uchoice_storage"):
        db.execute(text(f"delete from {table} where warehouse_code in (:a, :b)"), {"a": WH, "b": DEST})
    db.commit()


@pytest.fixture
def db():
    session = SessionLocal()
    _cleanup(session)
    yield session
    _cleanup(session)
    session.close()


def _stamp(db, actor, when):
    """Backdate every change written by `actor` (a unique created_by)."""
    db.execute(text("update uchoice_storage_txn set created_at = :ts where created_by = :actor"),
               {"ts": when, "actor": actor})
    db.commit()


def _delta(db, wh, sku, bpp, delta, txn_type, when, actor):
    from core.uchoice_storage import apply_storage_delta
    apply_storage_delta(db, wh, sku, bpp, delta, txn_type, None, note=None, created_by=actor)
    db.commit()
    _stamp(db, actor, when)


def _balances(db, warehouses=(WH,)):
    from core.uchoice_inventory import inventory_balances
    return {(b["warehouse_code"], b["sku_code"]): b
            for b in inventory_balances(db, list(warehouses), START, END_EXCLUSIVE)}


def _at(day, hour=12, **kw):
    return datetime.datetime(2031, 3, day, hour, tzinfo=UTC, **kw)


def test_each_change_type_lands_in_its_column(db):
    from core.uchoice_storage import apply_loose_pick

    _delta(db, WH, SKU, 20, 5, "inbound", datetime.datetime(2031, 2, 10, tzinfo=UTC), "inv-pre")
    _delta(db, WH, SKU, 20, 3, "opening", _at(1, 0), "inv-open")          # exactly at start
    _delta(db, WH, SKU, 20, 4, "inbound", _at(3), "inv-in")
    _delta(db, WH, SKU, 20, -2, "outbound", _at(5), "inv-out")
    _delta(db, WH, SKU, 20, -1, "adjust", _at(7), "inv-adj")
    _delta(db, WH, SKU, 20, -1, "transfer_out", _at(9), "inv-xfer")
    # Loose pick: 5 boxes from a 20-box pallet -> repackaging pair, no pallet movement
    apply_loose_pick(db, WH, SKU, 20, 5, None, "inv-pick")
    db.commit()
    _stamp(db, "inv-pick", _at(11))

    b = _balances(db)[(WH, SKU)]
    assert b["opening"] == 5 + 3          # before start + 'opening' inside the period
    assert b["inbound"] == 4
    assert b["outbound"] == -2            # the loose pick adds nothing here
    assert b["other"] == -1 - 1           # adjust + transfer_out; the pick pair nets 0
    assert b["closing"] == 8 + 4 - 2 - 2
    assert b["closing_mix"] == [(15, 1), (20, 7)]


def test_period_boundaries(db):
    _delta(db, WH, SKU, 20, 2, "inbound", datetime.datetime(2031, 2, 28, 23, 59, 59, 999999, tzinfo=UTC), "inv-b1")
    _delta(db, WH, SKU, 20, 3, "inbound", _at(1, 0), "inv-b2")
    _delta(db, WH, SKU, 20, 7, "inbound", datetime.datetime(2031, 4, 1, tzinfo=UTC), "inv-b3")

    b = _balances(db)[(WH, SKU)]
    assert b["opening"] == 2              # 1µs before start
    assert b["inbound"] == 3              # exactly at start is a movement
    assert b["closing"] == 5              # exactly at end_exclusive is next period


def test_inactive_zero_sku_excluded_but_netting_activity_kept(db):
    from core.uchoice_storage import apply_loose_pick

    _delta(db, WH, SKU_B, 10, 2, "inbound", datetime.datetime(2031, 1, 5, tzinfo=UTC), "inv-z1")
    _delta(db, WH, SKU_B, 10, -2, "outbound", datetime.datetime(2031, 1, 6, tzinfo=UTC), "inv-z2")
    _delta(db, WH, SKU, 20, 1, "inbound", datetime.datetime(2031, 2, 1, tzinfo=UTC), "inv-z3")
    apply_loose_pick(db, WH, SKU, 20, 5, None, "inv-z4")   # nets 0 pallets, but is activity
    db.commit()
    _stamp(db, "inv-z4", _at(2))

    balances = _balances(db)
    assert (WH, SKU_B) not in balances
    assert balances[(WH, SKU)]["closing"] == 1


def test_closing_equals_current_stock_when_period_ends_today(db):
    from core.uchoice_inventory import inventory_balances
    from core.uchoice_storage import apply_loose_pick, apply_storage_delta

    apply_storage_delta(db, WH, SKU, 20, 6, "inbound", None, note=None, created_by="inv-now")
    apply_storage_delta(db, WH, SKU, 20, -1, "outbound", None, note=None, created_by="inv-now")
    apply_loose_pick(db, WH, SKU, 20, 4, None, "inv-now")
    db.commit()

    today = datetime.date.today()
    [b] = inventory_balances(db, [WH], today.replace(day=1), today + datetime.timedelta(days=1))
    stock = db.execute(text(
        "select boxes_per_pallet, pallet_count from uchoice_storage "
        "where warehouse_code = :wh and pallet_count <> 0 order by 1"), {"wh": WH}).all()
    assert b["closing"] == sum(count for _, count in stock)
    assert b["closing_mix"] == [tuple(r) for r in stock]


def test_inventory_matches_completion_results(db):
    """Quantities, not sheets: the Inbound/Outbound columns equal the
    pallet_count the completion handlers report for the same requests."""
    from handlers.uchoice.storage_txns import ApplyInboundStorageHandler, ApplyOutboundStorageHandler

    inbound = ApplyInboundStorageHandler().handle({
        "wechat_openid": "inv-hin", "collected_fields": {},
        "_uchoice_target": {"warehouse_code": WH, "original_fields": {"sku_lines": [
            {"sku_code": SKU, "boxes_per_pallet": 20, "pallet_count": 6},
            {"sku_code": SKU_B, "boxes_per_pallet": 36, "pallet_count": 2},
        ]}},
    }, {}, db)
    db.commit()
    _stamp(db, "inv-hin", _at(4))

    outbound = ApplyOutboundStorageHandler().handle({
        "wechat_openid": "inv-hout", "collected_fields": {},
        "_uchoice_target": {"warehouse_code": WH, "original_fields": {"sku_lines": [
            {"sku_code": SKU, "boxes_per_pallet": 20, "pallet_count": 4},
        ]}},
    }, {}, db)
    db.commit()
    _stamp(db, "inv-hout", _at(6))

    balances = _balances(db)
    for line in inbound["received_lines"]:
        assert balances[(WH, line["sku_code"])]["inbound"] == line["pallet_count"]
    for line in outbound["fulfillment_lines"]:
        assert balances[(WH, line["sku_code"])]["outbound"] == -line["pallet_count"]


def test_transfer_is_other_on_both_sides(db):
    _delta(db, WH, SKU, 20, 5, "inbound", datetime.datetime(2031, 2, 1, tzinfo=UTC), "inv-t0")
    _delta(db, WH, SKU, 20, -2, "transfer_out", _at(8), "inv-t1")
    _delta(db, DEST, SKU, 20, 2, "transfer_in", _at(8), "inv-t2")

    balances = _balances(db, (WH, DEST))
    assert balances[(WH, SKU)]["other"] == -2 and balances[(WH, SKU)]["outbound"] == 0
    assert balances[(DEST, SKU)]["other"] == 2 and balances[(DEST, SKU)]["inbound"] == 0


def test_workbook_inventory_sheet(db):
    """The sheet follows Summary, carries a Closing formula, and lists the
    closing balance per pallet size one per line, as the chat reply does."""
    import core.uchoice_invoice_export as export
    from core.uchoice_storage import apply_loose_pick

    _delta(db, WH, SKU, 20, 5, "opening", _at(1, 1), "inv-w1")
    _delta(db, WH, SKU, 20, -1, "outbound", _at(2), "inv-w2")
    apply_loose_pick(db, WH, SKU, 20, 5, None, "inv-w3")
    db.commit()
    _stamp(db, "inv-w3", _at(3))

    data = export.build_invoice_workbook(
        db, WH, "2031-03", generated_at=datetime.datetime(2031, 4, 1, tzinfo=UTC))
    wb = load_workbook(io.BytesIO(data))
    assert wb.sheetnames == ["Summary", "Inventory", "Outbound", "Inbound", "Storage"]

    ws = wb["Inventory"]
    assert [c.value for c in ws[1]] == [
        "Warehouse", "SKU", "Description", "Opening (plt)", "Closing (plt)",
        "Inbound (plt)", "Outbound (plt)", "Other Net Change (plt)", "Closing Detail",
    ]
    row = [c.value for c in ws[2]]
    assert row[:2] == [WH, SKU]
    assert row[3] == 5 and row[5] == 0 and row[6] == -1 and row[7] == 0
    assert row[4] == "=D2+F2+G2+H2"
    assert row[8] == "1 托 @ 15/托\n3 托 @ 20/托"
    assert ws.cell(row=2, column=9).alignment.wrap_text
    assert [c.value for c in ws[3]][:1] == ["Total"]

    summary = {r[0]: r[1] for r in wb["Summary"].iter_rows(values_only=True) if r and r[0]}
    assert summary["Opening pallets"] == 5
    assert summary["Closing pallets"] == 4


def test_opening_label_shows_as_期初():
    from core import result_message, uchoice_storage_history_export
    assert result_message._TXN_TYPE_LABELS["opening"] == "期初"
    assert uchoice_storage_history_export._TXN_TYPE_LABELS["opening"] == "期初"
